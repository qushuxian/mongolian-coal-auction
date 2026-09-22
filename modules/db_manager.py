"""
Mongolian Coal Auction - 数据库管理与持久化模块 (db_manager.py)
----------------------------------------------------------------
功能定位:
1. 管理 SQLite (mongolian_coal_auction.db) 的初始化、建表与索引优化
2. 支持历史 Excel (蒙煤拍卖临时记录.xlsx) 一键清洗与增量归档迁移
3. 供 run_daily_sync.py 每日调用，实现当日爬取结果的幂等去重入库
4. 提供多维 SQL 查询、统计以及按需逆向导出 Excel 报表与前端 JSON
"""

import os
import re
import json
import sqlite3
from copy import copy
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(CURRENT_DIR)
DATA_DIR = os.path.join(PROJECT_ROOT, "data")
os.makedirs(DATA_DIR, exist_ok=True)

DEFAULT_DB_PATH = os.path.join(DATA_DIR, "mongolian_coal_auction.db")
DEFAULT_EXCEL_PATH = os.path.join(DATA_DIR, "蒙煤拍卖临时记录.xlsx")


def get_connection(db_path: str = None) -> sqlite3.Connection:
    """获取 SQLite 数据库连接，采用默认单文件日志模式，确保无临时缓存文件残留"""
    path = db_path or DEFAULT_DB_PATH
    conn = sqlite3.connect(path)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode = DELETE;")
    conn.execute("PRAGMA synchronous = NORMAL;")
    return conn


def init_db(db_path: str = None) -> None:
    """初始化数据库表结构与索引"""
    conn = get_connection(db_path)
    with conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS coal_auctions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                auction_date TEXT NOT NULL,             -- 拍卖日期 YYYY-MM-DD
                auction_time TEXT NOT NULL DEFAULT '10:00', -- 拍卖时间 HH:MM
                product TEXT NOT NULL DEFAULT '煤炭',   -- 产品类别
                coal_type TEXT NOT NULL,                -- 细分煤种 (如 洗精主焦煤 / 中挥发份性焦煤)
                seller TEXT NOT NULL,                   -- 卖方
                status TEXT,                            -- 状态 (成交 / 流拍 / 尚未开始 / 通告已发布 等)
                quantity_raw TEXT,                      -- 数量原始文本 (如 '2 批量/12800吨/')
                tons REAL,                              -- 提取出的纯吨数 (如 12800.0)
                lots INTEGER,                           -- 提取出的批次数 (如 2)
                currency TEXT NOT NULL DEFAULT 'CNY',   -- 币种 (CNY / USD)
                floor_price REAL NOT NULL,              -- 拍卖最低底价
                deal_price REAL,                        -- 成交价格每吨 (流拍为 0, 未开始为 NULL)
                price_change_rate REAL,                 -- 涨跌幅率 (如 0.340 代表 34.0%)
                price_type TEXT,                        -- 价格类型 (固定价格等)
                delivery_time TEXT,                     -- 供货时间
                delivery_location TEXT,                 -- 供货地点
                transport_mode TEXT,                    -- 运输方式
                pdf_url TEXT,                           -- 详情通告或报告 PDF 链接
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                UNIQUE(auction_date, auction_time, seller, coal_type, floor_price, quantity_raw)
            );
        """)

        # 常用时序与分析查询索引
        conn.execute("CREATE INDEX IF NOT EXISTS idx_auction_date ON coal_auctions(auction_date);")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_seller_coal ON coal_auctions(seller, coal_type);")
        conn.execute("CREATE INDEX IF NOT EXISTS idx_status ON coal_auctions(status);")
    conn.close()


def parse_quantity(quantity_str: str) -> Tuple[Optional[float], Optional[int]]:
    """从诸如 '10 批量/64000吨/' 或 '64000' 中提取 (总吨数, 批次数)"""
    if not quantity_str:
        return None, None
    raw = str(quantity_str)
    tons = None
    lots = None
    
    # 提取吨数
    ton_match = re.search(r'([\d\.,]+)\s*吨', raw)
    if ton_match:
        try:
            tons = float(ton_match.group(1).replace(',', ''))
        except ValueError:
            pass
    elif re.match(r'^\d+(\.\d+)?$', raw.strip()):
        try:
            tons = float(raw.strip())
        except ValueError:
            pass

    # 提取批次数
    lot_match = re.search(r'([\d]+)\s*批', raw)
    if lot_match:
        try:
            lots = int(lot_match.group(1))
        except ValueError:
            pass

    return tons, lots


def normalize_currency(currency_val: str, price_str: str = "") -> str:
    """标准化币种代码 (CNY / USD)"""
    combined = f"{currency_val or ''} {price_str or ''}".upper()
    if "USD" in combined or "$" in combined:
        return "USD"
    return "CNY"


def normalize_num(val: Any) -> Optional[float]:
    """数值安全提取与转换"""
    if val is None:
        return None
    if isinstance(val, (int, float)):
        return float(val)
    raw = str(val).strip()
    if not raw or raw in ["-", "null", "None", "/"]:
        return None
    nums = re.findall(r'[\d\.]+', raw.replace(',', ''))
    if nums:
        try:
            return float(nums[0])
        except ValueError:
            return None
    return None


def derive_status(status_raw: str, deal_price: Optional[float], raw_deal_str: str = "") -> str:
    """根据状态文本和成交价推导统一的状态标签"""
    status_text = str(status_raw or "").strip()
    deal_str = str(raw_deal_str or "").strip()
    
    if "流拍" in status_text or "流拍" in deal_str or "未提交报价" in deal_str:
        return "流拍"
    if deal_price is not None and deal_price > 0:
        return "已成交"
    if status_text in ["通告已发布", "尚未开始", "正在进行", "已结束"]:
        return status_text
    if deal_price == 0:
        return "流拍"
    return status_text or "未知"


def save_daily_auctions(query_date: str, items: List[Dict[str, Any]], db_path: str = None) -> int:
    """
    供 run_daily_sync.py 调用的主增量落库接口
    执行 INSERT OR REPLACE 幂等保存当日拍卖数据
    """
    init_db(db_path)
    conn = get_connection(db_path)
    saved_count = 0

    with conn:
        # 先清除当日可能残留的旧快照/未匹配通告，确保全量幂等覆盖
        conn.execute("DELETE FROM coal_auctions WHERE auction_date = ?", (query_date,))
        for item in items:
            # 基础属性
            prod = item.get("产品", "煤炭") or "煤炭"
            coal_type = item.get("产品类型", "") or ""
            seller = item.get("卖方", "") or ""
            raw_time = str(item.get("拍卖时间", "10:00")).strip()
            time_display = raw_time if raw_time else "10:00"

            # 数量
            qty_raw = item.get("数量", "") or ""
            tons, lots = parse_quantity(qty_raw)

            # 币种 & 价格
            raw_floor = item.get("拍卖最低价", "")
            raw_deal = item.get("成交价格每吨", "")
            curr = normalize_currency(item.get("币种", ""), f"{raw_floor} {raw_deal}")

            floor_p = normalize_num(raw_floor) or 0.0
            
            deal_p = None
            clean_deal = str(raw_deal).strip()
            if "流拍" in clean_deal or "未提交报价" in clean_deal or clean_deal in ["0", "0.00"]:
                deal_p = 0.0
            elif clean_deal and clean_deal not in ["-", "null", "None"]:
                deal_p = normalize_num(clean_deal)

            # 状态判定
            status = derive_status(item.get("状态", ""), deal_p, raw_deal)

            # 涨幅走势
            rate = None
            raw_rate = item.get("拍卖价格涨幅走势", "")
            if raw_rate and "%" in str(raw_rate):
                try:
                    rate = float(str(raw_rate).replace("%", "").strip()) / 100.0
                except ValueError:
                    rate = None
            elif deal_p is not None and floor_p > 0:
                rate = round((deal_p - floor_p) / floor_p, 4)

            # 条款与元数据
            price_type = item.get("价格类型", "")
            delivery_time = item.get("供货时间", "")
            delivery_loc = item.get("供货地点", "")
            transport = item.get("运输方式", "")
            pdf_url = item.get("查看具体信息", "")

            conn.execute("""
                INSERT OR REPLACE INTO coal_auctions (
                    auction_date, auction_time, product, coal_type, seller,
                    status, quantity_raw, tons, lots, currency,
                    floor_price, deal_price, price_change_rate, price_type,
                    delivery_time, delivery_location, transport_mode, pdf_url,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            """, (
                query_date, time_display, prod, coal_type, seller,
                status, qty_raw, tons, lots, curr,
                floor_p, deal_p, rate, price_type,
                delivery_time, delivery_loc, transport, pdf_url
            ))
            saved_count += 1

    conn.close()
    return saved_count


def migrate_from_excel(excel_path: str = None, db_path: str = None) -> int:
    """
    一键迁移历史 Excel (蒙煤拍卖临时记录.xlsx) 至 SQLite
    """
    import openpyxl

    path = excel_path or DEFAULT_EXCEL_PATH
    if not os.path.exists(path):
        raise FileNotFoundError(f"未找到历史 Excel 文件: {path}")

    init_db(db_path)
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = wb.active

    migrated_count = 0
    items_to_save = []

    for r in range(2, sheet.max_row + 1):
        product = sheet.cell(r, 1).value
        coal_type = sheet.cell(r, 2).value
        seller = sheet.cell(r, 3).value
        date_val = sheet.cell(r, 4).value
        qty_raw = sheet.cell(r, 5).value
        currency_raw = sheet.cell(r, 6).value
        floor_price_raw = sheet.cell(r, 7).value
        deal_price_raw = sheet.cell(r, 8).value
        # sheet.cell(r, 9) 为公式 (涨幅)
        delivery_time = sheet.cell(r, 10).value
        delivery_loc = sheet.cell(r, 11).value
        transport = sheet.cell(r, 12).value

        # 跳过完全空行
        if not coal_type and not seller and not date_val:
            continue

        # 解析日期与时间
        auction_date = ""
        auction_time = "10:00"
        if isinstance(date_val, datetime):
            auction_date = date_val.strftime("%Y-%m-%d")
            auction_time = date_val.strftime("%H:%M")
        elif date_val:
            dt_str = str(date_val).strip()
            if " " in dt_str:
                parts = dt_str.split(" ")
                auction_date = parts[0]
                auction_time = parts[1][:5]
            else:
                auction_date = dt_str[:10]

        tons, lots = parse_quantity(str(qty_raw or ""))
        curr = normalize_currency(str(currency_raw or ""), "")
        floor_p = normalize_num(floor_price_raw) or 0.0
        deal_p = normalize_num(deal_price_raw)

        # 判定状态
        status = "已成交" if (deal_p is not None and deal_p > 0) else ("流拍" if deal_p == 0 else "已发布")

        # 涨跌幅
        rate = None
        if deal_p is not None and floor_p > 0:
            rate = round((deal_p - floor_p) / floor_p, 4)

        items_to_save.append({
            "auction_date": auction_date,
            "auction_time": auction_time,
            "product": product or "煤炭",
            "coal_type": coal_type or "",
            "seller": seller or "",
            "status": status,
            "quantity_raw": str(qty_raw or ""),
            "tons": tons,
            "lots": lots,
            "currency": curr,
            "floor_price": floor_p,
            "deal_price": deal_p,
            "price_change_rate": rate,
            "delivery_time": str(delivery_time or ""),
            "delivery_location": str(delivery_loc or ""),
            "transport_mode": str(transport or "")
        })

    conn = get_connection(db_path)
    with conn:
        for it in items_to_save:
            conn.execute("""
                INSERT OR REPLACE INTO coal_auctions (
                    auction_date, auction_time, product, coal_type, seller,
                    status, quantity_raw, tons, lots, currency,
                    floor_price, deal_price, price_change_rate,
                    delivery_time, delivery_location, transport_mode,
                    updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
            """, (
                it["auction_date"], it["auction_time"], it["product"], it["coal_type"], it["seller"],
                it["status"], it["quantity_raw"], it["tons"], it["lots"], it["currency"],
                it["floor_price"], it["deal_price"], it["price_change_rate"],
                it["delivery_time"], it["delivery_location"], it["transport_mode"]
            ))
            migrated_count += 1
    conn.close()

    return migrated_count


def export_to_excel(output_path: str = None, db_path: str = None) -> str:
    """
    从 SQLite 逆向导出高兼容度的《蒙煤拍卖临时记录.xlsx》，完全保真原有 12 列排版与公式
    """
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    target_path = output_path or DEFAULT_EXCEL_PATH
    conn = get_connection(db_path)
    rows = conn.execute("""
        SELECT * FROM coal_auctions 
        ORDER BY auction_date ASC, auction_time ASC, id ASC
    """).fetchall()
    conn.close()

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Sheet1"

    headers = [
        '产品', '产品类型', '卖方', '日期', '数量',
        '币种', '拍卖最低价', '成交价格每吨', '涨幅走势',
        '供货时间', '供货地点', '运输方式'
    ]

    # 表头样式
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    for col_idx, h in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for r_idx, row in enumerate(rows, start=2):
        # 组装日期时间对象
        dt_val = row["auction_date"]
        try:
            y, m, d = map(int, row["auction_date"].split("-"))
            t_parts = row["auction_time"].split(":")
            hh, mm = int(t_parts[0]), int(t_parts[1])
            dt_obj = datetime(y, m, d, hh, mm)
        except Exception:
            dt_obj = f"{row['auction_date']} {row['auction_time']}"

        # 涨幅公式
        formula = f"=(H{r_idx}-G{r_idx})/G{r_idx}"

        row_vals = [
            row["product"],
            row["coal_type"],
            row["seller"],
            dt_obj,
            row["quantity_raw"],
            row["currency"],
            row["floor_price"],
            row["deal_price"],
            formula,
            row["delivery_time"],
            row["delivery_location"],
            row["transport_mode"]
        ]

        for col_idx, val in enumerate(row_vals, start=1):
            c = sheet.cell(row=r_idx, column=col_idx, value=val)
            c.font = Font(name="Microsoft YaHei", size=10)
            c.alignment = align_center if col_idx in [1, 2, 4, 6, 7, 8, 9] else align_left

            if col_idx == 4:
                c.number_format = 'yyyy-mm-dd hh:mm'
            elif col_idx in [7, 8]:
                c.number_format = '#,##0.00' if isinstance(val, float) and not val.is_integer() else 'General'
            elif col_idx == 9:
                c.number_format = '0.0%'

    # 列宽适配
    sheet.column_dimensions['A'].width = 10
    sheet.column_dimensions['B'].width = 16
    sheet.column_dimensions['C'].width = 28
    sheet.column_dimensions['D'].width = 18
    sheet.column_dimensions['E'].width = 18
    sheet.column_dimensions['F'].width = 8
    sheet.column_dimensions['G'].width = 14
    sheet.column_dimensions['H'].width = 14
    sheet.column_dimensions['I'].width = 12
    sheet.column_dimensions['J'].width = 32
    sheet.column_dimensions['K'].width = 36
    sheet.column_dimensions['L'].width = 14

    wb.save(target_path)
    return target_path


def get_summary_stats(db_path: str = None) -> Dict[str, Any]:
    """获取数据库全局统计指标"""
    init_db(db_path)
    conn = get_connection(db_path)
    
    total_records = conn.execute("SELECT COUNT(*) FROM coal_auctions").fetchone()[0]
    date_range = conn.execute("SELECT MIN(auction_date), MAX(auction_date) FROM coal_auctions").fetchone()
    deal_count = conn.execute("SELECT COUNT(*) FROM coal_auctions WHERE deal_price > 0").fetchone()[0]
    total_tons = conn.execute("SELECT SUM(tons) FROM coal_auctions").fetchone()[0] or 0.0

    coal_types = conn.execute("""
        SELECT coal_type, COUNT(*), SUM(tons) 
        FROM coal_auctions 
        GROUP BY coal_type 
        ORDER BY COUNT(*) DESC
    """).fetchall()

    sellers = conn.execute("""
        SELECT seller, COUNT(*), SUM(tons) 
        FROM coal_auctions 
        GROUP BY seller 
        ORDER BY COUNT(*) DESC
    """).fetchall()

    conn.close()
    return {
        "total_records": total_records,
        "earliest_date": date_range[0] if date_range else None,
        "latest_date": date_range[1] if date_range else None,
        "deal_count": deal_count,
        "total_tons": total_tons,
        "coal_types": [dict(ct) for ct in coal_types],
        "sellers": [dict(s) for s in sellers]
    }


def export_frontend_data(output_path: str = None, db_path: str = None) -> str:
    """
    从 SQLite 导出供前端 HTML 无需 CORS、零延迟直读的 auction_data.js
    挂载至全局 window.COAL_AUCTION_DATA
    """
    path = output_path or os.path.join(DATA_DIR, "auction_data.js")
    conn = get_connection(db_path)
    
    rows = conn.execute("""
        SELECT * FROM coal_auctions 
        ORDER BY auction_date ASC, auction_time ASC, id ASC
    """).fetchall()
    
    # 尝试获取今日 JSON 结构
    today_json_path = os.path.join(DATA_DIR, "daily_auction_merged.json")
    today_data = None
    if os.path.exists(today_json_path):
        try:
            with open(today_json_path, "r", encoding="utf-8") as f:
                today_data = json.load(f)
        except Exception:
            pass

    records = []
    for r in rows:
        deal_val = r["deal_price"]
        if deal_val is None:
            deal_display = 0 if r["status"] == "流拍" else "-"
        elif deal_val == 0:
            deal_display = 0
        else:
            deal_display = deal_val

        records.append({
            "产品": r["product"],
            "产品类型": r["coal_type"],
            "卖方": r["seller"],
            "日期": f"{r['auction_date']} {r['auction_time']}",
            "数量": r["quantity_raw"],
            "币种": r["currency"],
            "拍卖最低价": r["floor_price"],
            "成交价格每吨": deal_display,
            "供货时间": r["delivery_time"],
            "供货地点": r["delivery_location"],
            "运输方式": r["transport_mode"],
            "状态": r["status"],
            "查看具体信息": r["pdf_url"] or ""
        })

    conn.close()

    payload = {
        "update_time": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "total_records": len(records),
        "today": today_data,
        "records": records
    }

    content = f"// 蒙煤拍卖数据前端静态数据源 (由 db_manager.py 自动生成，支持 file:// 协议零拦截直接加载)\nwindow.COAL_AUCTION_DATA = {json.dumps(payload, ensure_ascii=False, indent=2)};\n"

    with open(path, "w", encoding="utf-8") as f:
        f.write(content)

    return path


if __name__ == "__main__":
    print("🚀 开始初始化 SQLite 数据库并执行历史数据迁移...")
    init_db()
    count = migrate_from_excel()
    print(f"✅ 成功从 Excel 迁移归档 {count} 条历史拍卖记录至 SQLite 数据库！")
    
    js_path = export_frontend_data()
    print(f"✅ 成功从 SQLite 导出前端免跨域数据源: {js_path}")
    
    stats = get_summary_stats()
    print("\n📊 数据库概览指标:")
    print(f"   • 总竞拍记录数: {stats['total_records']} 条")
    print(f"   • 日期范围跨度: {stats['earliest_date']} 至 {stats['latest_date']}")
    print(f"   • 累计成交场次: {stats['deal_count']} 场")
    print(f"   • 累计供标总量: {stats['total_tons']:,.0f} 吨")
    print("   • 煤种分布:")
    for ct in stats["coal_types"]:
        print(f"     - {ct['coal_type']}: {ct['COUNT(*)']} 场 ({ct['SUM(tons)'] or 0:,.0f} 吨)")
