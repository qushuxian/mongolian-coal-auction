#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Run Daily Sync & Merge (comex.mse.mn)
全自动获取当天实时拍卖看板与通告清单，进行数据关联（Join），
最终输出的 JSON 仅保留前端页面上实际展示的字段信息。
"""

import os
import re
import sys
import glob
import json
import argparse
from datetime import datetime

# 导入 modules 二级目录下的采集与数据管理模块
current_dir = os.path.dirname(os.path.abspath(__file__))
modules_dir = os.path.join(current_dir, "modules")
data_dir = os.path.join(current_dir, "data")
os.makedirs(data_dir, exist_ok=True)

if modules_dir not in sys.path:
    sys.path.insert(0, modules_dir)

from fetch_comex_auction import get_today_auctions, format_status
from fetch_notices import get_today_notices
import db_manager

def normalize_seller(name):
    if not name:
        return ""
    if "塔万陶勒盖" in name or "Tavantolgoi" in name or "Тавантолгой" in name:
        return "塔万陶勒盖"
    if "能源资源" in name or "Energy Resources" in name or "Энержи" in name:
        return "能源资源"
    if "蒙古罗斯" in name or "Монголросцветмет" in name or "Erdenes Critical" in name:
        return "蒙古罗斯"
    return re.sub(r'[^\w]', '', name)

def normalize_price(val):
    if val is None:
        return None
    nums = re.findall(r'[\d\.]+', str(val))
    if nums:
        try:
            # 当存在改价情况（如 '197.5$/181.5$/'）时，取修改后的最新底价 nums[-1]
            return float(nums[-1])
        except ValueError:
            pass
    return None

def normalize_time(time_str):
    if not time_str:
        return ""
    match = re.search(r'(\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2})', time_str)
    return match.group(1) if match else time_str.strip()

def format_currency_display(price, currency):
    """
    格式化为前端页面展示的价格格式，如 180.00 USD / 1,150.00 CNY
    """
    if price is None:
        return "-"
    curr_str = currency.upper().strip() if currency else ""
    return f"{float(price):,.2f} {curr_str}".strip()

def calculate_price_trend(deal_price, floor_price, is_unsold=False):
    """
    计算拍卖价格涨幅走势，如 0% 或 +5.2%，流拍则返回流拍
    """
    if is_unsold:
        return "流拍"
    if deal_price is None or floor_price is None or floor_price == 0:
        return "0%"
    diff = (deal_price - floor_price) / floor_price * 100
    if abs(diff) < 0.001:
        return "0%"
    return f"{diff:+.1f}%"

import io
import requests
import pdfplumber

pdf_delivery_cache = {}

def extract_pdf_delivery_info(pdf_url):
    """
    从《矿产品远期合同交易公告》PDF 中提取供货时间、供货地点、运输方式
    """
    if not pdf_url:
        return {"供货时间": "-", "供货地点": "-", "运输方式": "-"}
    if pdf_url in pdf_delivery_cache:
        return pdf_delivery_cache[pdf_url]

    res = {"供货时间": "-", "供货地点": "-", "运输方式": "-"}
    try:
        resp = requests.get(pdf_url, headers={"User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7)"}, timeout=15)
        if resp.status_code == 200:
            with pdfplumber.open(io.BytesIO(resp.content)) as pdf:
                if pdf.pages:
                    tables = pdf.pages[0].extract_tables()
                    if tables:
                        for row in tables[0]:
                            clean = [c for c in row if c is not None and str(c).strip() != ""]
                            if len(clean) >= 2:
                                key_col = str(clean[1])
                                val_col = str(clean[2]) if len(clean) >= 3 else str(clean[-1])
                                clean_val = re.sub(r'\s+', ' ', val_col).strip()
                                # 去除中文汉字之间因 PDF 换行产生的多余空格，保留英文/数字之间的空格
                                clean_val = re.sub(r'(?<=[\u4e00-\u9fa5])\s+(?=[\u4e00-\u9fa5])', '', clean_val)
                                if "供货时间" in key_col:
                                    res["供货时间"] = clean_val
                                elif "供货地点" in key_col:
                                    res["供货地点"] = clean_val
                                elif "运输方式" in key_col:
                                    res["运输方式"] = clean_val
    except Exception:
        pass

    pdf_delivery_cache[pdf_url] = res
    return res

EXCLUDED_PRODUCT_TYPES = {
    "fe-52%",
    "动力煤",
    "动力煤 /烟煤, 不粘结煤/",
    "washed non-coking coal"
}

def is_excluded_product_type(prod_type):
    """
    判断是否属于需要排除的产品类型（如 Fe-52%、动力煤 /烟煤, 不粘结煤/）
    """
    if not prod_type:
        return False
    normalized = str(prod_type).strip().lower()
    if normalized in EXCLUDED_PRODUCT_TYPES:
        return True
    return "fe-52%" in normalized or "动力煤" in normalized or "不粘结煤" in normalized or "washed non-coking coal" in normalized

def merge_frontend_records(auctions, notices):
    """
    根据时间、卖方、底价等维度将实时看板数据与通告数据关联，
    仅保留前端页面实际呈现的字段，并补充提取 PDF 中的交割条款。
    """
    merged_list = []
    used_notice_indices = set()

    for a in auctions:
        a_time = normalize_time(a.get("auctionStartTime", ""))
        a_seller = normalize_seller(a.get("sellerNameCN") or a.get("sellerNameMN"))
        a_price = normalize_price(a.get("productPrice"))

        matched_notice = None
        matched_idx = -1

        for idx, n in enumerate(notices):
            if idx in used_notice_indices:
                continue
            n_time = normalize_time(n.get("datetime", ""))
            n_seller = normalize_seller(n.get("seller", ""))
            n_price = normalize_price(n.get("floorPrice"))

            if a_time and n_time and a_time == n_time and a_seller and n_seller and a_seller == n_seller:
                if a_price is not None and n_price is not None and abs(a_price - n_price) < 0.01:
                    matched_notice = n
                    matched_idx = idx
                    break
                elif matched_notice is None:
                    matched_notice = n
                    matched_idx = idx

        if matched_notice:
            used_notice_indices.add(matched_idx)

        prod_type = matched_notice.get("productType") if matched_notice else "-"
        # 排除指定产品类型（如 Fe-52%）
        if is_excluded_product_type(prod_type):
            continue

        # 仅保留前端页面展示的字段信息
        currency = a.get("currency") or ("USD" if "$" in (matched_notice.get("floorPrice", "") if matched_notice else "") else "CNY")
        floor_num = a.get("productPrice")
        deal_num = a.get("orderPrice")

        # 格式化时间为前端页面展示形式（HH:MM）
        time_display = a_time.split(" ")[-1] if " " in a_time else a_time
        notice_pdf_url = matched_notice.get("detailInfoUrl") if matched_notice else None

        # 从公告 PDF 提取交割信息
        delivery = extract_pdf_delivery_info(notice_pdf_url)

        status_code = a.get("auctionStatus")
        # 官方 comex 页面中，当 auctionStatus == 3 且无成交价 (orderPrice 为空) 时，
        # 页面《成交价格每吨》显示为“买方未提交报价”，即代表流拍；或状态码 4 为流拍
        is_unsold = (status_code == 3 and (deal_num is None or deal_num == 0 or deal_num == "")) or (status_code == 4)

        status_display = "流拍" if is_unsold else format_status(status_code, deal_num)
        if deal_num is not None and deal_num > 0:
            deal_display = format_currency_display(deal_num, currency)
        elif is_unsold:
            deal_display = "流拍"
        else:
            deal_display = "-"

        item = {
            "拍卖时间": time_display,
            "产品": a.get("productTypeNameCN") or (matched_notice.get("product") if matched_notice else "煤炭"),
            "产品类型": prod_type,
            "卖方": a.get("sellerNameCN") or (matched_notice.get("seller") if matched_notice else "-"),
            "状态": status_display,
            "数量": matched_notice.get("quantity") if matched_notice else (f"{a.get('size')} 批量/{int(a.get('size',0)*float(a.get('lot_price',0)))}吨/" if a.get("lot_price") else "-"),
            "拍卖最低价": format_currency_display(floor_num, currency),
            "成交价格每吨": deal_display,
            "拍卖价格涨幅走势": calculate_price_trend(deal_num, floor_num, is_unsold=is_unsold),
            "价格类型": matched_notice.get("priceType") if matched_notice else "固定价格",
            "供货时间": delivery.get("供货时间", "-"),
            "供货地点": delivery.get("供货地点", "-"),
            "运输方式": delivery.get("运输方式", "-"),
            "查看具体信息": notice_pdf_url
        }
        merged_list.append(item)

    # 补充未在看板中但存在通告的记录
    for idx, n in enumerate(notices):
        if idx not in used_notice_indices:
            prod_type = n.get("productType")
            # 排除指定产品类型（如 Fe-52%）
            if is_excluded_product_type(prod_type):
                continue

            n_curr = "USD" if "$" in n.get("floorPrice", "") else ("CNY" if "¥" in n.get("floorPrice", "") else "")
            n_price = normalize_price(n.get("floorPrice"))
            time_raw = n.get("datetime", "")
            time_display = time_raw.split(" ")[-1][:5] if " " in time_raw else time_raw
            notice_pdf_url = n.get("detailInfoUrl")
            delivery = extract_pdf_delivery_info(notice_pdf_url)

            merged_list.append({
                "拍卖时间": time_display,
                "产品": n.get("product"),
                "产品类型": prod_type,
                "卖方": n.get("seller"),
                "状态": "通告已发布",
                "数量": n.get("quantity"),
                "拍卖最低价": format_currency_display(n_price, n_curr),
                "成交价格每吨": "-",
                "拍卖价格涨幅走势": "0%",
                "价格类型": n.get("priceType"),
                "供货时间": delivery.get("供货时间", "-"),
                "供货地点": delivery.get("供货地点", "-"),
                "运输方式": delivery.get("运输方式", "-"),
                "查看具体信息": notice_pdf_url
            })

    # 双重保险：过滤掉排除的产品类型
    merged_list = [item for item in merged_list if not is_excluded_product_type(item.get("产品类型"))]

    # 按拍卖时间升序排序
    merged_list.sort(key=lambda x: x.get("拍卖时间") or "")
    return merged_list

def clean_redundant_jsons(save_target):
    """
    自动清理 data 目录下所有其他的非目标 JSON 文件
    """
    all_jsons = glob.glob(os.path.join(data_dir, "*.json"))
    target_abs = os.path.abspath(save_target)
    for f in all_jsons:
        if os.path.abspath(f) != target_abs:
            try:
                os.remove(f)
            except Exception:
                pass

def print_frontend_table(query_date, items):
    print("\n" + "=" * 165)
    print(f"  蒙古煤炭拍卖综合数据清单 (含公告 PDF 供货条款)  |  日期: {query_date}  |  总场次: {len(items)}")
    print("=" * 165)

    if not items:
        print(f"  当前日期 ({query_date}) 暂无拍卖数据。")
        print("=" * 165)
        return

    header = f"{'时间':<6} | {'产品':<5} | {'产品类型':<16} | {'卖方':<22} | {'最低价':<13} | {'成交价':<13} | {'供货时间':<22} | {'运输方式':<8} | {'供货地点'}"
    print(header)
    print("-" * 165)

    for r in items:
        t = r.get("拍卖时间") or "-"
        p = r.get("产品") or "-"
        ptype = r.get("产品类型") or "-"
        seller = r.get("卖方") or "-"
        fprice = r.get("拍卖最低价") or "-"
        dprice = r.get("成交价格每吨") or "-"
        dtime = r.get("供货时间") or "-"
        dtype = r.get("运输方式") or "-"
        dloc = r.get("供货地点") or "-"

        print(f"{t:<6} | {p:<5} | {ptype:<16} | {seller:<22} | {fprice:<13} | {dprice:<13} | {dtime:<22} | {dtype:<8} | {dloc}")

    print("=" * 165 + "\n")

def sync_records_to_excel(query_date, frontend_records):
    """
    每次运行时，将 Excel 中当天 (query_date) 的旧数据删除，
    并把 JSON 中对应 Excel 字段的数据规范写入到 Excel 中。
    Excel 列结构：
    ['产品', '产品类型', '卖方', '日期', '数量', '币种', '拍卖最低价', '成交价格每吨', '涨幅走势', '供货时间', '供货地点', '运输方式']
    """
    excel_path = os.path.join(data_dir, "蒙煤拍卖临时记录.xlsx")
    if not os.path.exists(excel_path):
        print(f"⚠️ 未找到 Excel 文件: {excel_path}，跳过 Excel 同步。")
        return

    try:
        import openpyxl
        from copy import copy
    except ImportError:
        print("⚠️ 未安装 openpyxl 库，跳过 Excel 同步。")
        return

    print(f"\n[4/4] 正在将最新拍卖数据同步至 Excel 文件 ({os.path.basename(excel_path)})...")
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active

    # 1. 倒序检查并删除当天的旧记录
    deleted_count = 0
    for row_idx in range(sheet.max_row, 1, -1):
        date_cell = sheet.cell(row=row_idx, column=4)
        date_val = date_cell.value
        row_date_str = ""
        if isinstance(date_val, datetime):
            row_date_str = date_val.strftime("%Y-%m-%d")
        elif hasattr(date_val, "strftime"):
            row_date_str = date_val.strftime("%Y-%m-%d")
        elif isinstance(date_val, str):
            row_date_str = date_val.strip()[:10]
        
        if row_date_str == query_date:
            sheet.delete_rows(row_idx, 1)
            deleted_count += 1

    if deleted_count > 0:
        print(f"      -> 已删除 Excel 中原 [{query_date}] 当天的 {deleted_count} 行旧数据")
    else:
        print(f"      -> Excel 中暂无 [{query_date}] 当天历史数据，无需删除")

    # 2. 准备参考样式模板（取最后一行或第二行样式）
    ref_row = max(2, sheet.max_row)

    # 3. 追加写入当天最新数据
    added_count = 0
    for item in frontend_records:
        new_row_idx = sheet.max_row + 1

        # 产品
        product = item.get("产品", "煤炭") or "煤炭"
        # 产品类型
        prod_type = item.get("产品类型", "") or ""
        # 卖方
        seller = item.get("卖方", "") or ""

        # 日期 (datetime 对象)
        time_str = str(item.get("拍卖时间", "10:00")).strip()
        try:
            hour, minute = 10, 0
            if ":" in time_str:
                parts = time_str.split(":")
                hour, minute = int(parts[0]), int(parts[1])
            y, m, d = map(int, query_date.split("-")[:3])
            dt_obj = datetime(y, m, d, hour, minute)
        except Exception:
            dt_obj = f"{query_date} {time_str}"

        # 数量
        amount = item.get("数量", "") or ""

        # 币种 & 拍卖最低价 & 成交价格每吨
        raw_floor = item.get("拍卖最低价", "")
        raw_deal = item.get("成交价格每吨", "")

        # 识别币种
        currency = "CNY"
        if "USD" in str(raw_floor).upper() or "USD" in str(raw_deal).upper() or "$" in str(raw_floor):
            currency = "USD"
        else:
            currency = "CNY"

        # 最低价数值
        floor_num = None
        if raw_floor:
            nums = re.findall(r'[\d\.]+', str(raw_floor).replace(',', ''))
            if nums:
                try:
                    floor_num = float(nums[0])
                    if floor_num.is_integer():
                        floor_num = int(floor_num)
                except ValueError:
                    floor_num = None

        # 成交价数值
        deal_num = None
        status = item.get("状态", "")
        clean_deal = str(raw_deal).strip()
        if "流拍" in status or "流拍" in clean_deal or "未提交报价" in clean_deal or clean_deal in ["0", "0.00"]:
            deal_num = 0
        elif clean_deal and clean_deal not in ["-", "null", "None"]:
            nums = re.findall(r'[\d\.]+', clean_deal.replace(',', ''))
            if nums:
                try:
                    deal_num = float(nums[0])
                    if deal_num.is_integer():
                        deal_num = int(deal_num)
                except ValueError:
                    deal_num = None

        # 涨幅走势公式
        rate_formula = f"=(H{new_row_idx}-G{new_row_idx})/G{new_row_idx}"

        # 供货条款
        delivery_time = item.get("供货时间", "") or ""
        delivery_loc = item.get("供货地点", "") or ""
        transport = item.get("运输方式", "公路运输") or "公路运输"

        row_data = [
            product,        # A: 产品
            prod_type,      # B: 产品类型
            seller,         # C: 卖方
            dt_obj,         # D: 日期
            amount,         # E: 数量
            currency,       # F: 币种
            floor_num,      # G: 拍卖最低价
            deal_num,       # H: 成交价格每吨
            rate_formula,   # I: 涨幅走势
            delivery_time,  # J: 供货时间
            delivery_loc,   # K: 供货地点
            transport       # L: 运输方式
        ]

        # 写入单元格并设定格式
        for col_idx, val in enumerate(row_data, start=1):
            cell = sheet.cell(row=new_row_idx, column=col_idx, value=val)
            ref_cell = sheet.cell(row=ref_row, column=col_idx)
            if ref_cell.font:
                cell.font = copy(ref_cell.font)
            if ref_cell.alignment:
                cell.alignment = copy(ref_cell.alignment)

            # 针对特定列设定专用格式
            if col_idx == 4: # 日期
                cell.number_format = 'yyyy-mm-dd hh:mm'
            elif col_idx in [7, 8]: # 最低价、成交价
                cell.number_format = 'General'
            elif col_idx == 9: # 涨幅公式
                cell.number_format = '0.0%'

        added_count += 1

    wb.save(excel_path)
    print(f"      -> 成功向 Excel 重新写入 [{query_date}] 最新拍卖数据 {added_count} 笔！")
    print(f"    📊 {excel_path}")

def main():
    today_str = datetime.now().strftime("%Y-%m-%d")

    parser = argparse.ArgumentParser(description="获取并合并蒙古煤炭拍卖数据（仅保留前端展示字段）")
    parser.add_argument("--date", type=str, default=today_str, help=f"目标日期 (默认今天: {today_str})")
    args = parser.parse_args()

    query_date = args.date
    print(f"\n🚀 开始全自动同步并关联 [{query_date}] 蒙古煤炭拍卖数据...")

    # 1. 获取两路数据源
    print("[1/3] 正在从 comex.mse.mn 获取实时拍卖看板数据...")
    auctions = get_today_auctions(query_date)
    print(f"      -> 获取到 {len(auctions)} 场实时看板记录")

    print("[2/3] 正在从 comex.mse.mn 获取通告参与清单（中文环境）...")
    notices = get_today_notices(query_date)
    print(f"      -> 获取到 {len(notices)} 笔通告记录")

    # 2. 进行数据 Join 关联（仅提取前端展示字段）
    print("[3/3] 正在进行双源数据智能 Join 关联，生成前端专属展示字段...")
    frontend_records = merge_frontend_records(auctions, notices)

    # 3. 构造输出 JSON 结构
    final_output = {
        "日期": query_date,
        "更新时间": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "总场次": len(frontend_records),
        "数据清单": frontend_records
    }

    # 4. 统一存储到 data 目录下的 JSON 文件中
    unified_json_path = os.path.join(data_dir, "daily_auction_merged.json")
    with open(unified_json_path, "w", encoding="utf-8") as f:
        json.dump(final_output, f, indent=2, ensure_ascii=False)

    # 5. 持久化至 SQLite 数据库 (核心数据持久化层)
    print(f"\n💾 正在将 [{query_date}] 拍卖数据同步持久化至 SQLite 数据库...")
    try:
        saved_db_count = db_manager.save_daily_auctions(query_date, frontend_records)
        print(f"    -> 成功幂等写入 SQLite 数据库 ({saved_db_count} 笔记录)！")
        print(f"    📁 {os.path.join(data_dir, 'mongolian_coal_auction.db')}")
        js_export_path = db_manager.export_frontend_data()
        print(f"    -> 成功自动生成前端直连数据源: {js_export_path}")
    except Exception as e:
        print(f"    ⚠️ 写入 SQLite / 导出前端数据源时发生异常: {e}")

    # 6. 清理多余的其他 JSON 文件
    clean_redundant_jsons(unified_json_path)

    print(f"\n✅ 数据已成功提取，保留前端页面展示字段并更新 JSON:")
    print(f"    📄 {unified_json_path}")
    print(f"    🧹 已清除所有其他冗余 JSON。")

    # 7. 将最新数据同步重写至 Excel 文件中 (保持旧版工具向下兼容)
    sync_records_to_excel(query_date, frontend_records)

    # 8. 打印汇总表
    print_frontend_table(query_date, frontend_records)

if __name__ == "__main__":
    main()
